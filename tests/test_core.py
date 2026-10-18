from pathlib import Path
from types import SimpleNamespace

import docx

from app import hebrew
from app.agent import Agent, echo_safe
from app.config import Config
from app.index import Index


def make_library(root: Path) -> Path:
    lib = root / "הנחיות"
    (lib / "תכנון דרכים").mkdir(parents=True)
    d = docx.Document()
    d.add_heading("פרק 3 - חתך רוחב", level=1)
    d.add_paragraph("רוחב נתיב תנועה בדרך עירונית ראשית יהיה 3.25 מטר לפחות.")
    d.add_paragraph("בדרכים מקומיות ניתן להפחית את רוחב הנתיבים ל-3.00 מ'.")
    d.add_heading("פרק 4 - מעגלי תנועה", level=1)
    d.add_paragraph("קוטר מעגל חיצוני מינימלי בכיכר עירונית הוא 26 מטר.")
    d.save(lib / "תכנון דרכים" / "הנחיות לתכנון רחובות בערים.docx")
    (lib / "תמ\"א 42.txt").write_text("לפי תמ\"א 42 נדרש מרחק ראות לעצירה של 65 מ' במהירות תכן 50 קמ\"ש.", encoding="utf-8")
    (lib / "~$temp.docx").write_text("lock file")
    return lib


def test_hebrew_prefix_and_finals():
    assert "נתיבימ" in hebrew.index_text("והנתיבים").split()
    assert hebrew.tokens("כיכרות ם") == ["כיכרות", "מ"]
    assert hebrew.tokens('תמ"א 42') == ["תמא", "42"]
    assert hebrew.tokens("טבלה4.2 ל40 V/C") == ["טבלה", "4.2", "ל", "40", "v", "c"]
    q = hebrew.fts_query("רוחב הנתיב")
    assert '"נתיב"*' in q and '"רוחב"' in q


def test_visual_order_fix():
    logical = "רוחב נתיב תנועה בדרך עירונית יהיה שלושה מטרים לפחות כאשר המהירות היא חמישים קמ\"ש בשעות"
    visual = "\n".join(line[::-1] for line in logical.splitlines())
    assert hebrew.fix_visual_order(visual) == logical
    assert hebrew.fix_visual_order(logical) == logical


def test_index_search_read_incremental(tmp_path):
    lib = make_library(tmp_path)
    idx = Index(tmp_path / "db.sqlite")
    exts = [".docx", ".txt", ".pdf"]
    s = idx.rebuild([lib], exts)
    assert s["added"] == 2 and s["failed"] == 0

    hits = idx.search("מה רוחב הנתיבים בעיר?")
    assert hits and "3.25" in hits[0]["snippet"]
    assert any("26" in h["snippet"] for h in idx.search("כיכר"))
    assert idx.search("תמ\"א 42")[0]["title"] == "תמ\"א 42"

    doc_id = hits[0]["doc_id"]
    doc = idx.read(doc_id, 1, 10)
    assert doc["pages"] and doc["pages"][0]["location"].startswith("קטע 1")
    assert idx.list_documents("רחובות")[0]["doc_id"] == doc_id

    assert idx.rebuild([lib], exts)["unchanged"] == 2
    (lib / "תמ\"א 42.txt").unlink()
    assert idx.rebuild([lib], exts)["removed"] == 1
    assert idx.status()["documents"] == 1


def test_agent_tools(tmp_path, monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "test")
    lib = make_library(tmp_path)
    idx = Index(tmp_path / "db.sqlite")
    idx.rebuild([lib], [".docx", ".txt"])
    agent = Agent(Config(local_folders=[lib]), idx)
    sources: dict = {}
    out, err = agent._run_tool("search_guidelines", {"query": "רוחב נתיב", "doc_id": None, "max_results": 3}, sources)
    assert not err and "3.25" in out
    doc_id = next(iter(sources))
    out, err = agent._run_tool("read_pages", {"doc_id": doc_id, "page_from": 1, "page_to": 99}, sources)
    assert not err and sources[doc_id]["read"]
    _, err = agent._run_tool("read_pages", {"doc_id": 999, "page_from": 1, "page_to": 1}, sources)
    assert err
    _, err = agent._run_tool("search_guidelines", {"query": " ", "doc_id": None, "max_results": 3}, sources)
    assert err
    payload = agent._sources_payload(sources)
    assert payload[0]["read"] and payload[0]["title"] == "הנחיות לתכנון רחובות בערים"


def test_echo_safe_after_fallback():
    b = lambda t: SimpleNamespace(type=t)
    content = [b("thinking"), b("text"), b("tool_use"), b("fallback"), b("thinking"), b("tool_use")]
    assert [x.type for x in echo_safe(content)] == ["text", "fallback", "thinking", "tool_use"]
    plain = [b("thinking"), b("tool_use")]
    assert [x.type for x in echo_safe(plain)] == ["thinking", "tool_use"]
